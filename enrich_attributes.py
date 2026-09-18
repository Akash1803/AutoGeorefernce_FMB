r"""
enrich_attributes.py
====================
Post-process the FMB_Vector outputs:

1. Add location attributes to EVERY polygon feature (district, district_code,
   taluk, taluk_code, village, village_code, survey_no, giscode, sheet_scale,
   source_pdf, sketch_id) so each GeoJSON is self-describing when loaded alone.
2. Rewrite each <survey>_parcels.csv with the enriched columns.
3. Build one combined attribute table across all sheets:
       FMB_Vector\all_parcels_attributes.csv
       FMB_Vector\all_parcels_attributes.xlsx  (sheet "Parcels" + "Per sketch")
4. Build one combined GeoPackage, one layer per sketch, for convenient loading
   in QGIS (coordinates stay in local sheet metres per sketch):
       FMB_Vector\FMB_Vector_all_sketches.gpkg

Safe to re-run; it only rewrites derived files.
"""
import os, sys, csv, json, glob, collections, re
D = r"D:\Projects\Tambaram_Chengalpattu"; OUT = os.path.join(D, "FMB_Vector")
TALUK = {"04": "Chengalpattu", "05": "Tambaram", "15": "Vandalur"}
VILL = {}
for r in csv.DictReader(open(os.path.join(D, "FMB_Sketches", "fmb_sketch_summary.csv"), encoding="utf-8")):
    VILL[r["village_code"]] = r["village"]

LOC = ["district", "district_code", "taluk", "taluk_code", "village", "village_code", "survey_no", "giscode", "sketch_id", "sheet_scale", "source_pdf"]
KEEP = ["poly_id", "plot_no", "numbered", "area_sqm", "area_are", "area_hect", "perimeter_m", "n_vertices", "n_holes",
        "review_needed", "n_labels_in_polygon", "alt_labels", "label_source"]
all_rows = []; per_sketch = []
files = sorted(glob.glob(os.path.join(OUT, "35_*", "*_parcels.geojson")))
for gj in files:
    code = os.path.basename(os.path.dirname(gj)); stem = os.path.basename(gj).replace("_parcels.geojson", "")
    survey = re.sub(r"_s\d+$", "", stem)                       # 287_s2000 -> 287
    _, t, v = code.split("_")
    g = json.load(open(gj, encoding="utf-8")); md = g.get("metadata", {})
    loc = {"district": "Chengalpattu", "district_code": 35, "taluk": TALUK.get(t, t), "taluk_code": int(t),
           "village": VILL.get(code, code), "village_code": code, "survey_no": survey,
           "giscode": "S35%s%s%s" % (t, v, survey), "sketch_id": "%s_%s" % (code, survey),
           "sheet_scale": md.get("sheet_scale"), "source_pdf": "FMB_Sketches\\%s\\%s.pdf" % (code, stem)}
    single = len(g["features"]) == 1
    for f in g["features"]:
        p = f["properties"]
        new = {k: loc[k] for k in LOC}
        new.update({k: p.get(k) for k in KEEP})
        if single and not new["plot_no"]:
            # an unsubdivided survey: the one polygon IS the survey number
            new["plot_no"] = survey; new["label_source"] = "whole_survey_no_subdivision"
        f["properties"] = new
        all_rows.append(new)
    g["metadata"]["attributes_note"] = "Location attributes (district/taluk/village codes, giscode) added per feature by enrich_attributes.py"
    json.dump(g, open(gj, "w", encoding="utf-8"), indent=1)
    with open(gj.replace("_parcels.geojson", "_parcels.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=LOC + KEEP); w.writeheader()
        for f in sorted(g["features"], key=lambda f: (f["properties"]["plot_no"] or "zzz")): w.writerow(f["properties"])
    per_sketch.append({"sketch_id": loc["sketch_id"], "village": loc["village"], "village_code": code, "taluk": loc["taluk"], "survey_no": survey,
                       "polygons": len(g["features"]), "numbered": sum(1 for f in g["features"] if f["properties"]["plot_no"]),
                       "review_needed": sum(1 for f in g["features"] if f["properties"]["review_needed"]),
                       "polygon_total_sqm": md.get("polygon_total_sqm"), "vector_area_sqm": md.get("stated_area_sqm"),
                       "max_node_shift_m": md.get("max_node_displacement_m"), "text_class": md.get("text_class"), "geojson": os.path.relpath(gj, D)})

with open(os.path.join(OUT, "all_parcels_attributes.csv"), "w", newline="", encoding="utf-8") as fh:
    w = csv.DictWriter(fh, fieldnames=LOC + KEEP); w.writeheader(); w.writerows(all_rows)
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Parcels"
    ws.append(LOC + KEEP)
    for r in all_rows: ws.append([r[k] for k in LOC + KEEP])
    ws2 = wb.create_sheet("Per sketch"); hdr = list(per_sketch[0].keys()); ws2.append(hdr)
    for r in per_sketch: ws2.append([r[k] for k in hdr])
    for s in (ws, ws2):
        for c in s[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
        s.freeze_panes = "A2"; s.auto_filter.ref = s.dimensions
    wb.save(os.path.join(OUT, "all_parcels_attributes.xlsx"))
except Exception as e:
    print("xlsx skipped:", e)

# combined GeoPackage, one layer per sketch
try:
    import geopandas as gpd
    from shapely.geometry import shape
    gp = os.path.join(OUT, "FMB_Vector_all_sketches.gpkg")
    if os.path.exists(gp): os.remove(gp)
    for gj in files:
        g = json.load(open(gj, encoding="utf-8"))
        if not g["features"]: continue
        gdf = gpd.GeoDataFrame([f["properties"] for f in g["features"]], geometry=[shape(f["geometry"]) for f in g["features"]], crs=None)
        gdf.to_file(gp, layer=g["features"][0]["properties"]["sketch_id"], driver="GPKG")
    print("combined GeoPackage written:", gp)
except Exception as e:
    print("gpkg skipped:", e)

print("sheets:", len(files), "| polygons:", len(all_rows), "| numbered:", sum(1 for r in all_rows if r["plot_no"]),
      "| review_needed:", sum(1 for r in all_rows if r["review_needed"]))
print("per village:", collections.Counter(r["village"] for r in all_rows).most_common())
